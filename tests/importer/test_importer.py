"""Import: guards before parsing, header mapping, row validation, preview -> grant -> commit."""

import io
import re
import struct
import tracemalloc
import zipfile

import openpyxl
import pytest

from execution_world import OPERATOR, OPERATOR_2, VIEWER, FailingAudit
from twotrack.core.audit import AuditStatus
from twotrack.core.entities import ImportRejected, ImportService, importer, parse_upload
from twotrack.core.entities.importer import _collect
from twotrack.core.guards import UploadLimits, check_upload

CSV = (b"Part Code,Name,Qty On Hand,Thread_Pitch,Colour\n"
       b"F-10,Washer,12,0.5,grey\n"
       b"F-11,Nut,abc,0.7,\n"
       b"F-12,Spacer,1,,\n"
       b"F-1,Hex bolt,41,1.25,\n")


def service(world, **kwargs):
    return ImportService(world.registry, world.store, kwargs.pop("audit", world.audit),
                         world.grants, clock=world.clock, **kwargs)


def xlsx(rows):
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


MiB = 1024 * 1024


def archive(parts, method=zipfile.ZIP_DEFLATED):
    """A zip with the two parts every workbook has, plus ``parts`` (name -> bytes)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", method) as zipped:
        zipped.writestr("[Content_Types].xml", "<Types/>")
        zipped.writestr("xl/workbook.xml", "<workbook/>")
        for name, data in parts.items():
            zipped.writestr(name, data)
    return buffer.getvalue()


def declare(raw, name, size):
    """``raw`` with its central directory claiming ``size`` bytes for the entry ``name``: the
    one number a zip-bomb author changes, since zipfile writes the true size itself."""
    data, at = bytearray(raw), 0
    while (at := data.find(b"PK\x01\x02", at)) != -1:
        length = struct.unpack_from("<H", data, at + 28)[0]
        if data[at + 46:at + 46 + length].decode() == name:
            struct.pack_into("<I", data, at + 24, size)
            return bytes(data)
        at += 4
    raise AssertionError(f"no entry {name!r}")


def data_offset(raw, name):
    """Where the (compressed) bytes of the entry ``name`` start."""
    entry = zipfile.ZipFile(io.BytesIO(raw)).getinfo(name)
    name_length, extra_length = struct.unpack_from("<HH", raw, entry.header_offset + 26)
    return entry.header_offset + 30 + name_length + extra_length


def rejected(code, fn, *args, **kwargs):
    with pytest.raises(ImportRejected) as caught:
        fn(*args, **kwargs)
    assert caught.value.code == code
    return caught.value


def test_preview_maps_aliases_reports_unmapped_and_validates_rows(world):
    preview = service(world).preview(OPERATOR, "fastener", "parts.csv", CSV)
    assert [(m.header, m.field_id) for m in preview.mapping] == [
        ("Part Code", "sku"), ("Name", "name"), ("Qty On Hand", "qty"),
        ("Thread_Pitch", "thread_pitch")]
    assert preview.unmapped == ["Colour"]
    assert (preview.valid_count, preview.error_count) == (2, 2)
    by_row = {r.row: r for r in preview.rows}
    assert by_row[3].errors == ["Quantity: must be a number"]
    assert by_row[3].values["qty"] == "abc"  # the bad cell is shown as written
    assert by_row[4].errors == ["Thread Pitch: required"]  # category attr is required
    assert [by_row[2].action, by_row[5].action] == ["new", "update"]
    assert preview.grant is not None and world.store.get("part", "F-10") is None


def test_commit_writes_valid_rows_with_category_and_audits(world):
    svc = service(world)
    preview = svc.preview(OPERATOR, "fastener", "parts.csv", CSV)
    result = svc.commit(OPERATOR, "fastener", "parts.csv", CSV,
                        expected_digest=preview.digest, grant_token=preview.grant.token)
    assert (result.written, result.skipped) == (2, 2)
    assert world.store.get("part", "F-10")["_entity"] == "fastener"
    assert world.store.get("part", "F-1")["qty"] == 41
    assert world.audit_statuses() == [AuditStatus.NEEDS_CONFIRMATION, AuditStatus.STARTED,
                                      AuditStatus.OK]
    rejected("grant_invalid", svc.commit, OPERATOR, "fastener", "parts.csv", CSV,
             expected_digest=preview.digest, grant_token=preview.grant.token)


def test_commit_refuses_changed_file(world):
    svc = service(world)
    preview = svc.preview(OPERATOR, "fastener", "parts.csv", CSV)
    changed = CSV.replace(b"F-10,Washer,12", b"F-10,Washer,13")
    rejected("digest_mismatch", svc.commit, OPERATOR, "fastener", "parts.csv", changed,
             expected_digest=preview.digest, grant_token=preview.grant.token)
    assert world.store.get("part", "F-10") is None
    assert world.audit_statuses()[-1] is AuditStatus.BLOCKED


def test_commit_needs_the_uploaders_grant(world):
    svc = service(world)
    preview = svc.preview(OPERATOR, "fastener", "parts.csv", CSV)
    rejected("grant_invalid", svc.commit, OPERATOR_2, "fastener", "parts.csv", CSV,
             expected_digest=preview.digest, grant_token=preview.grant.token)


def test_viewer_cannot_import(world):
    rejected("forbidden", service(world).preview, VIEWER, "fastener", "parts.csv", CSV)
    assert world.audit_statuses() == [AuditStatus.BLOCKED]


def test_commit_with_failing_audit_writes_nothing(world):
    preview = service(world).preview(OPERATOR, "fastener", "parts.csv", CSV)
    failing = service(world, audit=FailingAudit(world.audit, {AuditStatus.OK}))
    with pytest.raises(Exception, match="disk full"):
        failing.commit(OPERATOR, "fastener", "parts.csv", CSV,
                       expected_digest=preview.digest, grant_token=preview.grant.token)
    assert world.store.get("part", "F-10") is None


def test_colliding_headers_reject_the_file(world):
    raw = b"Part Code,Item No,Name,Thread Pitch\nF-20,F-21,Bolt,1\n"
    error = rejected("header_collision", service(world).preview, OPERATOR, "fastener",
                     "x.csv", raw)
    assert "sku" in error.message
    # a sentence naming the columns and the field, never a Python list
    assert error.message == ("several columns map to the same field: “Part Code” and “Item No” "
                             "map to SKU (sku). Rename or remove the extra columns.")
    assert "[" not in error.message and "[" not in error.text("zh")
    assert error.text("zh").startswith("多个列对应到了同一个字段：")


def test_missing_required_column_rejects_the_file(world):
    error = rejected("missing_columns", service(world).preview, OPERATOR, "fastener", "x.csv",
                     b"Part Code,Name\nF-20,Bolt\n")
    assert error.message == "required columns not found: Thread Pitch"
    several = rejected("missing_columns", service(world).preview, OPERATOR, "fastener",
                       "x.csv", b"Name\nBolt\n")
    assert several.message == "required columns not found: SKU and Thread Pitch"
    assert several.text("zh") == "缺少必需的列：SKU 和 Thread Pitch"  # no Chinese label declared
    assert "[" not in several.message


def test_size_and_row_caps(world):
    small = service(world, limits=UploadLimits(max_bytes=64, max_rows=2))
    rejected("too_large", small.preview, OPERATOR, "fastener", "x.csv", CSV)
    rows = b"sku,thread_pitch\n" + b"".join(b"F-%d,1\n" % i for i in range(3))
    rejected("too_many_rows", small.preview, OPERATOR, "fastener", "x.csv", rows)
    rejected("empty", small.preview, OPERATOR, "fastener", "x.csv", b"")


def test_file_type_is_decided_by_magic_bytes(world):
    svc = service(world)
    for raw in (b"%PDF-1.7\n...", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", b"MZ\x90\x00",
                b"s\x00k\x00u\x00"):
        rejected("file_type", svc.preview, OPERATOR, "fastener", "parts.csv", raw)
    # the bytes decide, whatever the name: a workbook named .csv is still read as a workbook,
    # and a zip that is no workbook is refused
    assert svc.preview(OPERATOR, "fastener", "parts.csv", xlsx(
        [["SKU", "Thread Pitch"], ["F-60", 1.5]])).file_type == "xlsx"
    plain_zip = io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as archive:
        archive.writestr("notes.txt", "hello")
    rejected("file_type", svc.preview, OPERATOR, "fastener", "parts.xlsx", plain_zip.getvalue())


@pytest.mark.parametrize("name", ["parts.xlsx", "PARTS.XLSX", " parts.xlsx "])
def test_a_file_named_xlsx_that_is_no_workbook_is_a_file_type_problem(world, name):
    """A CSV (or a truncated download, or an HTML error page) saved as .xlsx used to be read
    as CSV and end in a baffling "required columns not found"."""
    svc = service(world)
    error = rejected("file_type", svc.preview, OPERATOR, "fastener", name, CSV)
    assert ".xlsx" in error.message and "workbook" in error.message
    assert ".xlsx" in error.text("zh") and "工作簿" in error.text("zh")
    assert world.audit_statuses() == [AuditStatus.INVALID]  # refused and on record
    rejected("file_type", svc.preview, OPERATOR, "fastener", name, b"<html>502 Bad Gateway</html>")
    # the name only ever speaks against the bytes: the same bytes under a CSV name are CSV
    assert svc.preview(OPERATOR, "fastener", "parts.csv", CSV).file_type == "csv"
    assert svc.preview(OPERATOR, "fastener", "parts", CSV).file_type == "csv"


def test_the_xlsx_name_check_leaves_other_refusals_their_own_message(world):
    svc = service(world)
    pdf = rejected("file_type", svc.preview, OPERATOR, "fastener", "parts.xlsx", b"%PDF-1.7\n...")
    assert "PDF" in pdf.message and ".xlsx" not in pdf.message
    rejected("empty", svc.preview, OPERATOR, "fastener", "parts.xlsx", b"")
    # the commit path re-reads the same bytes under the same name
    preview = svc.preview(OPERATOR, "fastener", "parts.csv", CSV)
    rejected("file_type", svc.commit, OPERATOR, "fastener", "parts.xlsx", CSV,
             expected_digest=preview.digest, grant_token=preview.grant.token)


def test_xlsx_zip_bomb_is_refused_before_parsing(world):
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", b" " * (30 * 1024 * 1024))
    raw = bomb.getvalue()
    assert len(raw) < 100_000  # tiny on the wire
    rejected("zip_bomb", service(world).preview, OPERATOR, "fastener", "b.xlsx", raw)
    ratio_only = service(world, limits=UploadLimits(max_uncompressed=10**9))
    rejected("zip_bomb", ratio_only.preview, OPERATOR, "fastener", "b.xlsx", raw)
    many = io.BytesIO()
    with zipfile.ZipFile(many, "w") as archive:
        for i in range(250):
            archive.writestr(f"part{i}.xml", "x")
    rejected("zip_bomb", service(world).preview, OPERATOR, "fastener", "m.xlsx",
             many.getvalue())


SHEET = "xl/worksheets/sheet1.xml"


def test_an_xlsx_part_that_holds_more_than_it_declares_is_refused_without_inflating_it(world):
    """The central directory is the archive's own claim; openpyxl inflates what is there."""
    raw = declare(archive({SHEET: bytes(16 * MiB)}), SHEET, 100)  # holds 16 MiB, says 100 bytes
    assert len(raw) < 100_000
    tracemalloc.start()
    try:
        rejected("zip_bomb", check_upload, raw, UploadLimits(), "b.xlsx")
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < 4 * MiB  # refused after a chunk or two, not after the 16 MiB
    rejected("zip_bomb", service(world).preview, OPERATOR, "fastener", "b.xlsx", raw)
    assert world.audit_statuses() == [AuditStatus.INVALID]


@pytest.mark.parametrize("method", ["bz2", "lzma"])
def test_xlsx_parts_may_only_be_stored_or_deflated(world, method):
    """Other methods are unbounded in older Pythons' zipfile, and no workbook writer uses them."""
    pytest.importorskip(method)
    compression = {"bz2": zipfile.ZIP_BZIP2, "lzma": zipfile.ZIP_LZMA}[method]
    honest = archive({SHEET: b"<sheetData/>"}, compression)
    rejected("zip_bomb", service(world).preview, OPERATOR, "fastener", "b.xlsx", honest)
    bomb = declare(archive({SHEET: bytes(4 * MiB)}, compression), SHEET, 100)
    rejected("zip_bomb", check_upload, bomb, UploadLimits(), "b.xlsx")


@pytest.mark.parametrize("method", [zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED])
def test_stored_and_deflated_workbooks_pass_the_archive_check(world, method):
    book = xlsx([["SKU", "Thread Pitch"], ["F-70", 1.5]])
    with zipfile.ZipFile(io.BytesIO(book)) as source:
        parts = {item.filename: source.read(item) for item in source.infolist()}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zipped:
        for name, data in parts.items():
            zipped.writestr(name, data, compress_type=method)
    preview = service(world).preview(OPERATOR, "fastener", "p.xlsx", out.getvalue())
    assert (preview.file_type, preview.valid_count) == ("xlsx", 1)


def test_an_xlsx_part_that_cannot_be_inflated_is_an_unreadable_workbook_not_a_crash(world):
    stored = archive({SHEET: b"<sheetData/>" * 50}, zipfile.ZIP_STORED)
    flipped = bytearray(stored)
    flipped[data_offset(stored, SHEET) + 3] ^= 0xFF  # the CRC no longer matches
    error = rejected("file_type", service(world).preview, OPERATOR, "fastener", "b.xlsx",
                     bytes(flipped))
    assert "not a readable XLSX workbook" in error.message
    deflated = archive({SHEET: b"<sheetData/>" * 500})
    junk = bytearray(deflated)
    at = data_offset(deflated, SHEET)
    junk[at:at + 4] = b"\xff\xff\xff\xff"  # not a deflate stream
    rejected("file_type", service(world).preview, OPERATOR, "fastener", "b.xlsx", bytes(junk))


def test_real_xlsx_is_parsed_with_types(world):
    raw = xlsx([["SKU", "Name", "Quantity", "Thread Pitch"], ["F-30", "Pin", 5, 0.4],
                [None, None, None, None], ["F-31", "Rivet", 6.0, 1]])
    preview = service(world).preview(OPERATOR, "fastener", "p.xlsx", raw)
    assert preview.file_type == "xlsx" and preview.valid_count == 2
    assert preview.rows[1].values == {"sku": "F-31", "name": "Rivet", "qty": 6,
                                      "thread_pitch": 1.0}


# ---------------------------------------------------------------- rows and columns read


def with_dimension(raw, ref):
    """``raw`` (a workbook) with the <dimension> of its sheet replaced by ``ref``: the one claim
    a file makes about its own size, and the one a reader is tempted to believe."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            part = source.read(item)
            if item.filename == SHEET:
                part = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="%s"' % ref.encode(),
                              part)
            target.writestr(item, part)
    return out.getvalue()


def sparse_xlsx(last_row):
    """A header in row 1 and one data row at ``last_row``, no row in between, under a
    <dimension> that says the sheet is as wide as xlsx goes (16384 columns)."""
    book = openpyxl.Workbook()
    book.active.append(["sku", "thread_pitch"])
    book.active.cell(row=last_row, column=1, value="F-80")
    book.active.cell(row=last_row, column=2, value=1.5)
    buffer = io.BytesIO()
    book.save(buffer)
    return with_dimension(buffer.getvalue(), f"A1:XFD{last_row}")


def test_blank_rows_count_toward_the_rows_read_cap():
    limits = UploadLimits(max_rows=5)
    assert limits.max_rows_read == 11  # the header, and a blank row after each of 5 data rows
    spaced = b"sku,thread_pitch\n" + b"F-1,1\n\n" * 5
    assert len(parse_upload(spaced, limits).rows) == 5
    error = rejected("too_many_rows", parse_upload, spaced + b"\n", limits)
    assert "11" in error.message and "11" in error.text("zh")
    # one data row is far below max_rows, but the blank lines around it are rows read too
    rejected("too_many_rows", parse_upload, b"sku,thread_pitch\n" + b"\n" * 100 + b"F-1,1\n",
             limits)


def test_the_reader_is_not_asked_for_more_rows_than_the_cap_allows():
    pulled = []

    def records():
        yield ["sku", "thread_pitch"]
        for row in range(100_000):  # a file that is nearly all empty rows
            pulled.append(row)
            yield []

    rejected("too_many_rows", _collect, records(), UploadLimits(max_rows=5))
    assert len(pulled) == 11  # it stops at the first row past the cap, not after counting all


def test_a_sparse_sheet_is_cut_off_at_the_rows_read_cap():
    """A <dimension> that claims A1:XFD40 made openpyxl hand over 40 rows of 16384 cells; the
    empty ones were skipped before they were counted, so the cap never tripped."""
    raw = sparse_xlsx(40)
    assert len(raw) < 10_000  # a few KB on the wire
    rejected("too_many_rows", parse_upload, raw, UploadLimits(max_rows=5), "s.xlsx")
    table = parse_upload(raw, UploadLimits(max_rows=20), "s.xlsx")  # 2 * 20 + 1 reaches row 40
    assert [number for number, _ in table.rows] == [40]


def test_a_sparse_sheet_is_read_no_wider_than_its_header(world):
    raw = sparse_xlsx(8)
    table = parse_upload(raw, UploadLimits(max_rows=5), "s.xlsx")
    assert table.headers == ["sku", "thread_pitch"]  # not 16384 names, the rest of them empty
    assert table.rows == [(8, ["F-80", 1.5])]  # not 16384 cells either
    preview = service(world).preview(OPERATOR, "fastener", "s.xlsx", raw)
    assert (preview.valid_count, preview.rows[0].row) == (1, 8)


def test_a_sheet_costs_what_it_holds_not_what_it_claims(monkeypatch):
    """The cells that reach the cell reader. The file holds 5; its <dimension> (A1:XFD5) and a
    cell in the last column of row 3 each offer thousands more, and a row may be looked at only
    as far as max_columns + 1. The padding hides well (empty header cells are trimmed, rows are
    cut to the header), so only counting shows whether it was ever walked."""
    book = openpyxl.Workbook()
    book.active.append(["sku", "thread_pitch"])
    book.active.cell(row=3, column=16384, value="far right")
    book.active.cell(row=5, column=1, value="F-80")
    book.active.cell(row=5, column=2, value=1.5)
    buffer = io.BytesIO()
    book.save(buffer)
    looked = []
    real = importer._cell
    monkeypatch.setattr(importer, "_cell", lambda cell: looked.append(cell) or real(cell))
    parse_upload(with_dimension(buffer.getvalue(), "A1:XFD5"), UploadLimits(max_columns=10),
                 "s.xlsx")
    assert len(looked) < 30  # 2 + 11 + 2 once the dimension is dropped and the row is sliced


def test_an_understated_dimension_does_not_hide_rows():
    """The sheet's own <dimension> is not believed in either direction: a claim of A1:B2 over
    four rows used to drop the last two without a word."""
    book = xlsx([["sku", "thread_pitch"], ["F-81", 1], ["F-82", 2], ["F-83", 3]])
    table = parse_upload(with_dimension(book, "A1:B2"), UploadLimits(), "u.xlsx")
    assert [number for number, _ in table.rows] == [2, 3, 4]


def test_a_data_row_is_kept_no_wider_than_the_header():
    table = parse_upload(b"sku,thread_pitch\nF-90,1,left,over\nF-91\n", UploadLimits())
    assert table.rows == [(2, ["F-90", "1"]), (3, ["F-91"])]


def test_a_header_wider_than_max_columns_is_refused():
    limits = UploadLimits(max_columns=4)
    wide = b"sku,thread_pitch,name,qty,status\nF-1,1,Pin,2,active\n"
    error = rejected("too_many_columns", parse_upload, wide, limits)
    assert "4" in error.message and "4" in error.text("zh")
    # empty header cells at the end are padding (a styled row ends in dozens), not columns
    padded = b"sku,thread_pitch,name,qty" + b"," * 50 + b"\nF-1,1,Pin,2\n"
    assert parse_upload(padded, limits).headers == ["sku", "thread_pitch", "name", "qty"]
    # an empty cell inside the header still counts as a column
    rejected("too_many_columns", parse_upload, b"sku,,thread_pitch,name,qty\nF-1,,1,Pin,2\n",
             limits)


def test_the_default_header_may_have_1024_columns_and_no_more():
    def header(count):
        return (",".join(f"c{i}" for i in range(count)) + "\n").encode()

    assert UploadLimits().max_columns == 1024
    assert len(parse_upload(header(1024), UploadLimits()).headers) == 1024
    rejected("too_many_columns", parse_upload, header(1025), UploadLimits())
    wide_book = xlsx([[f"c{i}" for i in range(1025)]])
    rejected("too_many_columns", parse_upload, wide_book, UploadLimits(), "w.xlsx")


def test_encoding_fallback():
    limits = UploadLimits()
    text = "sku,name\nF-40,螺栓\n"  # two CJK characters
    table = parse_upload(text.encode("gb18030"), limits)
    assert table.encoding == "gb18030" and table.rows[0][1][1] == "螺栓"
    table = parse_upload(b"\xef\xbb\xbf" + text.encode(), limits)
    assert table.encoding == "utf-8-sig" and table.headers == ["sku", "name"]
    table = parse_upload("sku,name\nF-41,Café,\n".encode("latin-1"), limits)
    assert table.encoding == "latin-1" and table.rows[0][1][1] == "Café"


def test_semicolon_csv_and_duplicate_keys(world):
    raw = b"sku;thread_pitch\nF-50;1\nF-50;2\n"
    preview = service(world).preview(OPERATOR, "fastener", "s.csv", raw)
    assert preview.valid_count == 1
    assert preview.rows[1].errors == ["duplicate key; first seen in row 2"]
